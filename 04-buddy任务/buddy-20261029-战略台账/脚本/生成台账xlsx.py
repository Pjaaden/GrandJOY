#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成《全国招商总部 · 运动/男装战略品牌合同台账》Excel 工作簿。

用法：
    python 生成台账xlsx.py --源文件 <BI导出的xlsx> [--输出 <路径>] [--说明 "备注"]
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from 台账核心 import SHEET_SPECS, build_all  # noqa: E402

# ── 视觉规范（浅色主题）────────────────────────────────────────────────────
C_HEAD = "1F3864"        # 表头深蓝
C_HEAD_FONT = "FFFFFF"
C_BAND = "F2F6FC"        # 隔行浅蓝
C_SUB = "E8EEF7"         # 次级表头
C_TOTAL = "FFF2CC"       # 合计行
C_WARN = "FDE9E9"        # 校验异常
C_WARN_FONT = "B42318"
C_OK_FONT = "067647"
C_BORDER = "D6DCE4"

THIN = Side(style="thin", color=C_BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FONT_TITLE = Font(name="微软雅黑", size=14, bold=True, color="1F3864")
FONT_HEAD = Font(name="微软雅黑", size=10, bold=True, color=C_HEAD_FONT)
FONT_BODY = Font(name="微软雅黑", size=10)
FONT_TOTAL = Font(name="微软雅黑", size=10, bold=True)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def _is_num_col(header: str) -> bool:
    return any(k in header for k in ("(元)", "面积", "租金", "扣率", "天数", "月数", "差异", "数", "编号")) and "期数" not in header or header in (
        "合同期间", "期次数", "合同数", "铺位行数", "序号",
    )


def write_sheet(wb, name: str, rows: list[dict], spec: dict, freeze: str = "A2"):
    """按统一样式写一个子表。"""
    ws = wb.create_sheet(name)
    headers = spec["headers"]
    dtypes = spec.get("dtypes", {})
    formats = spec.get("formats", {})

    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = FONT_HEAD
        cell.fill = PatternFill("solid", fgColor=C_HEAD)
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER
    ws.row_dimensions[1].height = 32

    for i, row in enumerate(rows):
        ws.append([row.get(h) for h in headers])
        r = i + 2
        band = (i % 2 == 1)
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=c)
            cell.font = FONT_BODY
            cell.border = BORDER
            dt_name = dtypes.get(h, "object")
            if dt_name.startswith(("int", "float")):
                cell.alignment = ALIGN_RIGHT
                cell.number_format = formats.get(h, "#,##0.00")
            else:
                cell.alignment = ALIGN_LEFT if h in ("合同名称", "商户名称", "备注", "口径与计算规则", "源表来源列", "字段名称") else ALIGN_CENTER
            if band:
                cell.fill = PatternFill("solid", fgColor=C_BAND)

    # 列宽
    for c, h in enumerate(headers, start=1):
        width = 12
        if h in ("合同名称", "商户名称", "备注", "口径与计算规则", "源表来源列", "字段名称", "项目名称"):
            width = 30
        elif h in ("数据校验", "涉及三级业态", "所属区域", "品牌/行业"):
            width = 18
        elif h in ("资源位编号", "合同编号", "三楼", "起租日", "到期日", "合同起租日", "合同到期日"):
            width = 16
        elif "(元)" in h:
            width = 18
        ws.column_dimensions[get_column_letter(c)].width = width

    # 冻结 + 筛选
    ws.freeze_panes = freeze
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    return ws


def add_dashboard(wb, data: dict, meta: dict):
    """首张「台账看板」：关键指标 + 区域分布 + 品牌 Top。"""
    ws = wb.create_sheet("台账看板", 0)
    detail = data["合同台账明细"]
    contract = data["合同级汇总"]
    region = data["区域业态汇总"]
    brand = data["品牌汇总"]

    total_area = sum(c["签约面积(平方米)"] or 0 for c in contract)
    total_rent = sum(c["期间固定租金总额(元)"] or 0 for c in contract)
    total_cost = sum(c["期间成本合计(元)"] or 0 for c in contract)

    ws["A1"] = "全国招商总部 · 运动/男装板块 战略品牌合同台账"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 28

    ws["A2"] = f"数据版本：{meta.get('stat_date', '—')}　|　源文件：{meta.get('file', '—')}　|　生成时间：{dt.datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(name="微软雅黑", size=9, color="666666")
    ws.merge_cells("A2:F2")

    # 关键指标
    kpis = [
        ("合同数（份）", len(contract)),
        ("明细行数（期次）", len(detail)),
        ("覆盖项目数", len({d["项目名称"] for d in detail if d["项目名称"]})),
        ("签约面积合计（平方米）", round(total_area, 2)),
        ("期间固定租金合计（元）", round(total_rent, 2)),
        ("期间成本合计（元）", round(total_cost, 2)),
    ]
    r0 = 4
    for i, (k, v) in enumerate(kpis):
        col = 1 + (i % 3) * 2
        row = r0 + (i // 3) * 2
        ws.cell(row=row, column=col, value=k).font = Font(name="微软雅黑", size=9, color="666666")
        c = ws.cell(row=row + 1, column=col, value=v)
        c.font = Font(name="微软雅黑", size=14, bold=True, color="1F3864")
        if isinstance(v, float):
            c.number_format = "#,##0.00"
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)

    # 区域分布
    row = r0 + 5
    ws.cell(row=row, column=1, value="一、区域分布（合同级口径）").font = Font(name="微软雅黑", size=11, bold=True, color="1F3864")
    row += 1
    hdr = ["所属区域", "合同数", "签约面积（㎡）", "期间固定租金（元）", "期间成本合计（元）", "成本占比"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = FONT_HEAD
        c.fill = PatternFill("solid", fgColor=C_SUB)
        c.alignment = ALIGN_CENTER
        c.border = BORDER
    region_agg = {}
    for c in contract:
        b = region_agg.setdefault(c["所属区域"] or "(未标注)", {"n": 0, "area": 0.0, "rent": 0.0, "cost": 0.0})
        b["n"] += 1
        b["area"] += c["签约面积(平方米)"] or 0
        b["rent"] += c["期间固定租金总额(元)"] or 0
        b["cost"] += c["期间成本合计(元)"] or 0
    for i, (k, v) in enumerate(sorted(region_agg.items(), key=lambda x: -x[1]["cost"])):
        row += 1
        vals = [k, v["n"], round(v["area"], 2), round(v["rent"], 2), round(v["cost"], 2),
                (v["cost"] / total_cost) if total_cost else 0]
        for j, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=j, value=val)
            c.font = FONT_BODY
            c.border = BORDER
            c.alignment = ALIGN_RIGHT if j > 1 else ALIGN_LEFT
            if j == 3:
                c.number_format = "#,##0.00"
            elif j in (4, 5):
                c.number_format = "#,##0.00"
            elif j == 6:
                c.number_format = "0.0%"
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=C_BAND)

    # 品牌 Top15
    row += 2
    ws.cell(row=row, column=1, value="二、品牌 Top15（按期间成本合计）").font = Font(name="微软雅黑", size=11, bold=True, color="1F3864")
    row += 1
    hdr2 = ["品牌/行业", "合同数", "项目数", "签约面积（㎡）", "期间成本合计（元）", "平均抽成扣率"]
    for j, h in enumerate(hdr2, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = FONT_HEAD
        c.fill = PatternFill("solid", fgColor=C_SUB)
        c.alignment = ALIGN_CENTER
        c.border = BORDER
    for i, b in enumerate(brand[:15]):
        row += 1
        vals = [b["品牌/行业"], b["合同数"], b["项目数"], b["签约面积合计(平方米)"],
                b["期间成本合计(元)"], b["平均抽成扣率"]]
        for j, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=j, value=val)
            c.font = FONT_BODY
            c.border = BORDER
            c.alignment = ALIGN_RIGHT if j > 1 else ALIGN_LEFT
            if j == 4:
                c.number_format = "#,##0.00"
            elif j == 5:
                c.number_format = "#,##0.00"
            elif j == 6 and val is not None:
                c.number_format = "0.00%"
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=C_BAND)

    # 使用说明
    row += 2
    ws.cell(row=row, column=1, value="三、台账使用与更新说明").font = Font(name="微软雅黑", size=11, bold=True, color="1F3864")
    notes = [
        "1. 明细粒度 = 合同期次（合同期间列：1=第一期，2=第二期……）。起租日/到期日为该期次边界，非整份合同边界。",
        "2. 面积类指标一律看「合同级汇总」与「区域业态汇总」，避免期次重复计数。",
        "3. 「数据校验」列标注每行与 BI 口径的一致性：一致 / 口径差异·两者取高 / 免租期超期·负值 / 抽成制·BI无固定租金。",
        "4. 抽成扣率以小数存储（0.16 = 16%），可直接参与计算；抽成制合同无固定租金，日租金与固定租金总额为空属正常。",
        "5. 更新方式：把新一版 BI 导出放入「源数据」目录，运行 脚本/更新台账.sh；本地 xlsx 重新生成并回写飞书同名表。",
        "6. 飞书位置：我的空间 / 大悦商业 / 战略台账 / 「全国运动男装战略品牌合同台账」。",
    ]
    for n in notes:
        row += 1
        c = ws.cell(row=row, column=1, value=n)
        c.font = Font(name="微软雅黑", size=9, color="333333")
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)

    for col, w in zip("ABCDEF", [30, 14, 18, 22, 22, 14]):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    return ws


def highlight_checks(wb):
    """把「数据校验」非「一致」的行标红提示。"""
    ws = wb["合同台账明细"]
    headers = [c.value for c in ws[1]]
    if "数据校验" not in headers:
        return
    idx = headers.index("数据校验") + 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=idx).value
        if v and v != "一致":
            cell = ws.cell(row=r, column=idx)
            cell.fill = PatternFill("solid", fgColor=C_WARN)
            cell.font = Font(name="微软雅黑", size=10, color=C_WARN_FONT)
            cell.alignment = ALIGN_CENTER


def main():
    ap = argparse.ArgumentParser(description="生成运动/男装战略品牌合同台账")
    ap.add_argument("--源文件", required=True, help="BI 导出的合同台账 xlsx")
    ap.add_argument("--输出", default=None, help="输出 xlsx 路径")
    ap.add_argument("--说明", default="", help="版本备注")
    args = ap.parse_args()

    data = build_all(args.源文件, note=args.说明)
    from 台账核心 import load_source
    _, meta = load_source(args.源文件)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    add_dashboard(wb, data, meta)
    for name, spec in SHEET_SPECS.items():
        write_sheet(wb, name, data[name], spec)
    highlight_checks(wb)

    out = args.输出 or os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "产出", f"全国运动男装战略品牌合同台账_{meta.get('stat_date', 'v1')}.xlsx",
    )
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)

    print(f"[OK] 已生成：{out}")
    for name in SHEET_SPECS:
        print(f"     - {name}: {len(data[name])} 行")
    print(f"     - 合同数 {len(data['合同级汇总'])} / 明细 {len(data['合同台账明细'])} 行")
    return out


if __name__ == "__main__":
    main()
